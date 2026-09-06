"""Build Valentin-Presentation-v5.pptx.

Takes v3, drops its five bullet-and-banner Deep-Dive slides (6-10, Identity
among them) and puts nine deck-v2-styled slides in their place.
"""
from pptx import Presentation
from pptx.util import Emu

ROOT = '/Users/coralst/Desktop/screenshots/projects/valentin-romantic-agent'
WT = ROOT + '/.claude/worktrees/deck-html-deepdive'
TMP = '/Users/coralst/.claude/jobs/84959c04/tmp'
V3 = ROOT + '/docs/Valentin-Presentation-v3.pptx'
V5 = WT + '/docs/Valentin-Presentation-v5.pptx'

# in presentation order; (png, alt text)
NEW = [
    ('v3-1-scorecard.png',    'Two engines, scored on four dimensions'),
    ('v3-2-runtime.png',      'Runtime - where the agent runs'),
    ('v3-3-memory.png',       'Memory - how the agent remembers'),
    ('v3-4-gateway.png',      'Gateway - how the agent uses tools'),
    ('v3-5-observability.png','Observability - what you can see'),
    ('good-cost.png',         'Two bills, split into what is fixed and what grows'),
    ('good-blast.png',        'One fault. Forty sessions.'),
    ('good-team.png',         "I didn't write it. They did."),
    ('good-graph.png',        'Every node is a real pull request'),
]
DROP = range(5, 10)   # 0-based: v3 slides 6..10
INSERT_AT = 5

# the one live link on these nine slides: the #graph CTA
EXPLORE = 'https://d26dwovftfq9oe.cloudfront.net/explore.html#work'
PX = 12192000 / 1600.0            # EMU per rendered pixel
CTA_BOX = (80, 591, 252, 37)      # measured from the live #graph CTA

prs = Presentation(V3)
assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)
sld_lst = prs.slides._sldIdLst
ids0 = list(sld_lst)

blank = min(prs.slide_layouts, key=lambda l: len(l.placeholders))
added = []
for png, alt in NEW:
    s = prs.slides.add_slide(blank)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    pic = s.shapes.add_picture(f'{TMP}/{png}', 0, 0,
                               width=prs.slide_width, height=prs.slide_height)
    pic._element.nvPicPr.cNvPr.set('descr', alt)
    added.append(s)

    if png == 'good-graph.png':
        x, y, w, h = (Emu(int(v * PX)) for v in CTA_BOX)
        box = s.shapes.add_textbox(x, y, w, h)
        box.fill.background()
        box.line.fill.background()
        run = box.text_frame.paragraphs[0].add_run()
        run.text = ' '
        run.hyperlink.address = EXPLORE

for i in sorted(DROP, reverse=True):
    el = ids0[i]
    prs.part.drop_rel(el.rId)
    sld_lst.remove(el)

# move the nine appended slides into position 5
ids = list(sld_lst)
new_els = ids[-len(NEW):]
for el in new_els:
    sld_lst.remove(el)
for off, el in enumerate(new_els):
    sld_lst.insert(INSERT_AT + off, el)

prs.save(V5)

check = Presentation(V5)
print('v5 slides:', len(check.slides._sldIdLst))
for i, s in enumerate(check.slides, 1):
    pics = [sh for sh in s.shapes if sh.shape_type == 13]
    txt = [sh.text_frame.text.split('\n')[0][:45] for sh in s.shapes
           if sh.has_text_frame and sh.text_frame.text.strip()]
    if pics and not txt:
        print(i, 'IMG', pics[0]._element.nvPicPr.cNvPr.get('descr'))
    else:
        print(i, txt[:2])
links = 0
for s in check.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.hyperlink.address:
                        links += 1
                        print('LINK', r.hyperlink.address)
print('hyperlinks:', links)
