"""build_arch.py — writes the architecture slide out as a one-slide pptx.

    python3 docs/slide-variants/build_arch.py

Deliberately standalone: the slide is meant to be looked at on its own and then
dropped into docs/Valentin-Presentation-v6.pptx in place of the flattened render
that is in there now. No badge and no page number, since neither is known until
it has a position in the deck.
"""
import os
import sys

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import arch_slide                                          # noqa: E402
import emit_pptx                                           # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, os.pardir, 'Valentin-Architecture-v6.pptx')


def build(out=OUT):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.3333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])      # blank
    emit_pptx.render(slide, arch_slide.ops())
    slide.notes_slide.notes_text_frame.text = arch_slide.NOTES
    prs.save(out)
    return os.path.abspath(out)


if __name__ == '__main__':
    print(build())
