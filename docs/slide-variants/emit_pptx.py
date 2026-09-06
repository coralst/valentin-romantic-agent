"""emit_pptx.py — renders aws_slides ops onto a python-pptx slide.

Native shapes, not a flattened picture: the text stays editable in PowerPoint,
which is the whole point of authoring these slides as ops.
"""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import lxml.etree as etree

import aws_slides as T

ALIGN = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
ANCHOR = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'


def _fill(shape, color):
    if color is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)


def _stroke(shape, line):
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()


def _add_tail_end(line):
    ln = line._get_or_add_ln()
    for old in ln.findall(qn('a:tailEnd')):
        ln.remove(old)
    ln.append(ln.makeelement(qn('a:tailEnd'),
                             {'type': 'triangle', 'w': 'med', 'len': 'med'}))


def _set_text(shape, runs, size, color, bold, align, valign):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[valign]
    p = tf.paragraphs[0]
    p.alignment = ALIGN[align]
    if isinstance(runs, str):
        runs = [(runs, bold, color)]
    for text, rbold, rcolor in runs:
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = T.FONT
        f.size = Pt(size)
        f.bold = bool(rbold)
        f.color.rgb = RGBColor.from_string(rcolor or color)


def render(slide, ops):
    for op in ops:
        kind = op[0]
        if kind == 'rect':
            _, x, y, w, h, fill, line, radius = op
            if radius:
                sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                            Inches(y), Inches(w), Inches(h))
                sh.adjustments[0] = min(0.5, radius / min(w, h))
            else:
                sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                            Inches(w), Inches(h))
            _fill(sh, fill)
            _stroke(sh, line)
            sh.shadow.inherit = False
        elif kind == 'oval':
            _, x, y, w, h, fill, line = op
            sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                        Inches(w), Inches(h))
            _fill(sh, fill)
            _stroke(sh, line)
            sh.shadow.inherit = False
        elif kind == 'text':
            _, x, y, w, h, runs, size, color, bold, align, valign = op
            box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            _set_text(box, runs, size, color, bold, align, valign)
        elif kind in ('arrow', 'line'):
            _, x1, y1, x2, y2, color, wpt = op
            cx = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                            Inches(y1), Inches(x2), Inches(y2))
            cx.line.color.rgb = RGBColor.from_string(color)
            cx.line.width = Pt(wpt)
            if kind == 'arrow':
                _add_tail_end(cx.line)
        elif kind == 'arc':
            _, x, y, w, h, start, swing, color, wpt, arrow = op
            add_arc(slide, x, y, w, h, start, start + swing, color, wpt, arrow)
        else:
            raise ValueError(kind)


ARC_XML = """<p:sp xmlns:p="{P}" xmlns:a="{A}">
  <p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
    <a:prstGeom prst="arc"><a:avLst>
      <a:gd name="adj1" fmla="val {adj1}"/><a:gd name="adj2" fmla="val {adj2}"/>
    </a:avLst></a:prstGeom>
    <a:noFill/>
    <a:ln w="{lw}" cap="rnd"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>
      {tail}</a:ln>
  </p:spPr>
  <p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>
</p:sp>"""


def add_arc(slide, x, y, w, h, start_deg, end_deg, color, wpt, arrow=True):
    """A preset 'arc' sweeps clockwise from adj1 to adj2; 0deg is 3 o'clock."""
    sid = 1000 + len(slide.shapes._spTree)
    xml = ARC_XML.format(
        P=P, A=A, sid=sid, name=f'Arc {sid}',
        x=int(Inches(x)), y=int(Inches(y)), cx=int(Inches(w)), cy=int(Inches(h)),
        adj1=int(round(start_deg * 60000)) % 21600000,
        adj2=int(round(end_deg * 60000)) % 21600000,
        lw=int(Pt(wpt)), color=color,
        tail='<a:tailEnd type="triangle" w="med" len="med"/>' if arrow else '')
    sp = etree.fromstring(xml)
    slide.shapes._spTree.append(sp)
    return sp
