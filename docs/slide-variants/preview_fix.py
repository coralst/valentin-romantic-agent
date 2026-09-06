"""preview_fix.py — HTML preview of the three de-rasterised slides.

The graph slide keeps its plot as a picture, so this extracts that picture out
of the deck and crops it the same way the build script does, letting the
preview show exactly what PowerPoint will render.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from PIL import Image

import emit_html
import raster_slides as R
import graph_crop as G
from graph_crop import CROP

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, '..', 'Valentin-Presentation-v7.pptx')
OUTDIR = os.path.join(HERE, 'v7-previews')
TOTAL = 17


def cropped_plot():
    """Pull graph.png out of the deck and crop it to the plot region."""
    prs = Presentation(DECK)
    for s in prs.slides:
        for sh in s.shapes:
            if sh.shape_type == 13 and sh.image.size == (2400, 1352) and sh.crop_left:
                blob = sh.image.blob
                break
    raw = os.path.join(OUTDIR, 'graph-raw.png')
    os.makedirs(OUTDIR, exist_ok=True)
    open(raw, 'wb').write(blob)
    im = Image.open(raw)
    w, h = im.size
    box = (int(CROP['l'] * w), int(CROP['t'] * h),
           int((1 - CROP['r']) * w), int((1 - CROP['b']) * h))
    out = os.path.join(OUTDIR, 'graph-plot.png')
    im.crop(box).save(out)
    print('plot crop', box, '->', im.crop(box).size)
    return 'v7-previews/graph-plot.png'


def main():
    plot = cropped_plot()
    slides = [R.slide_bill(8, TOTAL, 7), R.slide_team(14, TOTAL, 13),
              G.slide_graph(15, TOTAL, 14, plot)]
    out = os.path.join(HERE, 'preview-fix.html')
    emit_html.emit(slides, out, labels={0: 'bill', 1: 'team', 2: 'graph'})
    print(out)


if __name__ == '__main__':
    main()
