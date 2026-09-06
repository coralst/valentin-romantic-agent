"""preview_v8.py — HTML preview of every natively-authored v8 slide.

Pixel-comparable to PowerPoint at 1600px wide (120 dpi), so this is what the
deck check looks at instead of a deploy-and-squint. Regenerate after any ops
change; the file is the review artifact.

    python3 docs/slide-variants/preview_v8.py
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from PIL import Image

import emit_html
import lens_slides as L
import cost_chart as C
import raster_slides as R
import graph_crop as G

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, '..', 'Valentin-Presentation-v8.pptx')
OUTDIR = os.path.join(HERE, 'previews')
TOTAL = 17
PAGES = dict(cost=7, resil=9, security=11, ownership=12, debug=13)
BADGES = dict(cost=6, resil=8, security=10, ownership=11, debug=12)


def cropped_plot():
    """Pull the PR-graph screenshot out of the deck and crop it as the deck does.

    The preview then shows the same pixels PowerPoint will, so the native lane
    labels can be checked against the picture's bands.
    """
    blob = None
    for s in Presentation(DECK).slides:
        for sh in s.shapes:
            if sh.shape_type == 13 and sh.image.size == (2400, 1352):
                blob = sh.image.blob
    if blob is None:
        raise RuntimeError('no graph picture in the deck — build it first')

    os.makedirs(OUTDIR, exist_ok=True)
    im = Image.open(io.BytesIO(blob))
    w, h = im.size
    box = (int(G.CROP['l'] * w), int(G.CROP['t'] * h),
           int((1 - G.CROP['r']) * w), int((1 - G.CROP['b']) * h))
    im.crop(box).save(os.path.join(OUTDIR, 'graph-plot.png'))
    return 'previews/graph-plot.png'


def main():
    plot = cropped_plot()
    slides = [L.slide_matrix(6, TOTAL, 5),
              C.slide_cost_graph(PAGES['cost'], TOTAL, BADGES['cost']),
              R.slide_bill(8, TOTAL, 7)]
    slides += L.all_lens_slides(TOTAL, PAGES, BADGES)[:1]      # resilience
    slides += [L.slide_blast(10, TOTAL, 9)]
    slides += L.all_lens_slides(TOTAL, PAGES, BADGES)[1:]      # security → debug
    slides += [R.slide_team(14, TOTAL, 13),
               G.slide_graph(15, TOTAL, 14, plot)]
    labels = {0: 'matrix', 1: 'cost graph', 2: 'bill', 3: 'resilience',
              4: 'blast', 5: 'security', 6: 'ownership', 7: 'debuggability',
              8: 'team', 9: 'PR graph'}
    out = os.path.join(HERE, 'preview-v8.html')
    emit_html.emit(slides, out, labels=labels)
    print(out)


if __name__ == '__main__':
    main()
