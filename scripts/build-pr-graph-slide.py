#!/usr/bin/env python3
"""Rebuild the PR-graph slide with current numbers — same design, new data.

The slide is unchanged in layout: the lane labels stay live text down the left,
the graph is a picture in the same box, the four tiles and both notes sit where
they sat. Only the figures move, and they all come from
``scripts/deck_stats.py`` so this slide and the team slide cannot disagree.

The graph itself is re-rendered rather than reused, because the picture in the
deck still shows 57 PRs. The picture box has a fixed aspect, so the render
computes the node spacing that makes today's PR count fit it, and shrinks the
nodes to match — the PR numbers inside the nodes are dropped at that size
because they would be unreadable. Nothing else about the graph changes.

Run:  python3 scripts/build-pr-graph-slide.py
Out:  pptx/Valentin-Slide-PR-Graph.pptx  +  docs/assets/graph/pr-graph-slide.png
"""

from __future__ import annotations

import importlib.util
import pathlib
import subprocess
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import deck_stats  # noqa: E402

ROOT = pathlib.Path(__file__).resolve().parent.parent
GRAPH_DIR = ROOT / "docs" / "assets" / "graph"
SLIDE_SVG = GRAPH_DIR / "pr-graph-slide.svg"
SLIDE_PNG = GRAPH_DIR / "pr-graph-slide.png"
OUT = ROOT / "pptx" / "Valentin-Slide-PR-Graph.pptx"

# ── the slide's geometry, measured off the deck it replaces ─────────────────
PIC_L, PIC_T, PIC_W, PIC_H = 2.05, 2.35, 10.67, 2.4583
LANE_LABEL_L, LANE_LABEL_W = 0.62, 1.32
LANE_LABEL_T, LANE_LABEL_H, LANE_LABEL_STEP = 2.68, 0.26, 0.31
TILE_T, TILE_W, TILE_H, TILE_STEP = 5.42, 2.9, 0.86, 3.0625
TILE_L = 0.62

NAVY = RGBColor(0x23, 0x2F, 0x3E)
ORANGE = RGBColor(0xFF, 0x99, 0x00)
MUTED = RGBColor(0x54, 0x5B, 0x64)
CARD = RGBColor(0xF7, 0xF8, 0xFA)
BORDER = RGBColor(0xE4, 0xE7, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

RENDER_ZOOM = 3  # px per SVG user unit in the rasterised graph


def graph_module():
    spec = importlib.util.spec_from_file_location(
        "agent_graph", ROOT / "scripts" / "generate-agent-graph.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def render_graph(pr_count: int) -> None:
    """Render the graph so its lane rows land on the slide's lane labels.

    The crop is derived, not guessed: the label rows fix where lane 0 and lane 6
    must fall inside the picture box, which fixes the crop's height and top; the
    box's aspect ratio then fixes its width, and therefore the node spacing.
    """
    graph = graph_module()
    lane_span = (len(graph.AGENTS) - 1) * graph.LANE_H          # lane 0 → lane 6
    label_span = (len(graph.AGENTS) - 1) * LANE_LABEL_STEP
    units_per_inch = lane_span / label_span
    crop_h = PIC_H * units_per_inch
    first_label_centre = LANE_LABEL_T + LANE_LABEL_H / 2
    crop_y0 = graph.LANE_TOP - (first_label_centre - PIC_T) * units_per_inch
    crop_w = crop_h * (PIC_W / PIC_H)

    # Clear of the graph's own gutter labels — the slide puts those down the
    # left as live text, and the widest of them ("56 authored · 0 reviewed · …")
    # runs to about x=237.
    crop_x0 = graph.GUTTER - 10
    slot = (crop_w + crop_x0 - graph.GUTTER - graph.PAD_RIGHT) / pr_count
    node_scale = round(min(1.0, slot / graph.SLOT * 1.1), 3)

    subprocess.run(
        [sys.executable, str(ROOT / "scripts" / "generate-agent-graph.py"),
         "--slot", f"{slot:.4f}", "--node-scale", f"{node_scale}",
         "--no-node-labels", "--out", SLIDE_SVG.name],
        cwd=ROOT, check=True)

    full = GRAPH_DIR / "pr-graph-slide.full.png"
    subprocess.run(["rsvg-convert", "-z", str(RENDER_ZOOM),
                    str(SLIDE_SVG), "-o", str(full)], check=True)
    with Image.open(full) as img:
        box = (round(crop_x0 * RENDER_ZOOM), round(crop_y0 * RENDER_ZOOM),
               round((crop_x0 + crop_w) * RENDER_ZOOM),
               round((crop_y0 + crop_h) * RENDER_ZOOM))
        img.crop(box).save(SLIDE_PNG)
    full.unlink()
    print(f"  slot {slot:.2f}, nodes ×{node_scale}, "
          f"crop {box[2] - box[0]}×{box[3] - box[1]} px")


def text_box(slide, left, top, width, height, *, anchor=MSO_ANCHOR.TOP,
             align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.paragraphs[0].alignment = align
    return frame


def write(frame, text, *, size=11, colour=MUTED, bold=False, spacing=None):
    para = frame.paragraphs[0]
    if spacing is not None:
        para.line_spacing = spacing
    run = para.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour
    return para


def rect(slide, left, top, width, height, fill, *, shape=MSO_SHAPE.RECTANGLE,
         line=None, adjust=None):
    box = slide.shapes.add_shape(shape, Inches(left), Inches(top),
                                 Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = line
        box.line.width = Pt(1)
    box.shadow.inherit = False
    if adjust is not None and box.adjustments:
        box.adjustments[0] = adjust
    return box


def build(stats):
    render_graph(stats["prs"])

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])

    rect(slide, 0, 0, 13.333, 7.5, WHITE)
    badge = rect(slide, 0.55, 0.4, 0.65, 0.65, NAVY, shape=MSO_SHAPE.OVAL)
    badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    badge.text_frame.margin_left = badge.text_frame.margin_right = 0
    write(badge.text_frame, "15", size=26, colour=ORANGE, bold=True)

    write(text_box(slide, 1.4, 0.4, 11.0, 0.65, anchor=MSO_ANCHOR.MIDDLE),
          "Every node is a real pull request", size=30, colour=NAVY, bold=True)
    rect(slide, 0.55, 1.15, 1.5, 0.04, ORANGE)
    write(text_box(slide, 0.55, 1.4, 12.2, 0.42),
          f"All {stats['prs']} of them, in the order they were opened, connected "
          f"to where they merged.", size=14, colour=MUTED)
    write(text_box(slide, 0.55, 7.05, 9.0, 0.35),
          "VALENTIN   ·   GenAI TFC Capstone Deep-Dive",
          size=10, colour=MUTED, bold=True)
    write(text_box(slide, 11.3, 7.05, 1.5, 0.35, align=PP_ALIGN.RIGHT),
          "16 / 18", size=10, colour=MUTED)

    write(text_box(slide, 0.62, 1.96, 12.1, 0.28, anchor=MSO_ANCHOR.MIDDLE),
          "node size = files changed   ·   ring = a genuine multi-agent review "
          "thread   ·   dashed = reviewed, then closed without merging   ·   "
          "left to right = the order the PRs were opened",
          size=10.5, colour=MUTED)

    graph = graph_module()
    for i, (_, name, _, _) in enumerate(graph.AGENTS):
        frame = text_box(slide, LANE_LABEL_L, LANE_LABEL_T + i * LANE_LABEL_STEP,
                         LANE_LABEL_W, LANE_LABEL_H,
                         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
        write(frame, name, size=10, colour=NAVY, bold=True)

    slide.shapes.add_picture(str(SLIDE_PNG), Inches(PIC_L), Inches(PIC_T),
                             Inches(PIC_W), Inches(PIC_H))

    tiles = [
        (f"{stats['prs']}", "pull requests",
         f"{stats['merged']} merged · {stats['closed']} closed · "
         f"{stats['open']} open"),
        (f"{stats['commits']}", "commits",
         f"+{stats['added']:,} lines, {stats['files']:,} files"),
        (f"{stats['comments']}", "review comments",
         f"the longest thread ran {stats['longest_thread']} rounds"),
        (f"{stats['authored']['infra']}", "workflow PRs",
         "the methodology is the biggest lane"),
    ]
    for i, (value, label, sub) in enumerate(tiles):
        left = TILE_L + i * TILE_STEP
        rect(slide, left, TILE_T, TILE_W, TILE_H, CARD,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=BORDER, adjust=0.12)
        write(text_box(slide, left + 0.18, TILE_T + 0.06, 1.0, 0.4,
                       anchor=MSO_ANCHOR.MIDDLE),
              value, size=20, colour=NAVY, bold=True)
        write(text_box(slide, left + 1.16, TILE_T + 0.06, 1.56, 0.4,
                       anchor=MSO_ANCHOR.MIDDLE),
              label, size=12, colour=NAVY, bold=True)
        write(text_box(slide, left + 0.18, TILE_T + 0.48, 2.54, 0.32,
                       anchor=MSO_ANCHOR.MIDDLE),
              sub, size=10, colour=MUTED)

    write(text_box(slide, 0.62, 6.44, 12.1, 0.52),
          f"Generated from the repo’s PR history by a script, so it cannot drift. "
          f"Node size is files changed; a ring means a genuine multi-agent review "
          f"thread; the axis is PR order, not a calendar. At the peak, "
          f"{stats['peak_open']} PRs were open at once.",
          size=10.5, colour=MUTED, spacing=0.95)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}  "
          f"({stats['prs']} PRs, {stats['commits']} commits, "
          f"{stats['comments']} review comments)")


if __name__ == "__main__":
    build(deck_stats.load())
