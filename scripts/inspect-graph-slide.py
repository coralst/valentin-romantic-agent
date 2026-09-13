#!/usr/bin/env python3
"""Print the alignments and picture crop of the PR-graph slide, so a rebuild
can reproduce them exactly.

Usage:  python3 scripts/inspect-graph-slide.py <deck.pptx> <slide-number>
"""

import sys

from pptx import Presentation

deck = Presentation(sys.argv[1])
slide = deck.slides[int(sys.argv[2]) - 1]

for i, sh in enumerate(slide.shapes):
    if sh.has_text_frame and sh.text_frame.text.strip():
        aligns = [p.alignment for p in sh.text_frame.paragraphs]
        print(i, repr(sh.text_frame.text[:40]), "align", aligns,
              "anchor", sh.text_frame.vertical_anchor)
    if sh.shape_type == 13:
        print(i, "PICTURE crop l/r/t/b:",
              sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom)
        print("   native size", sh.image.size, "placed",
              sh.width / 914400, "x", sh.height / 914400)
