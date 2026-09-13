#!/usr/bin/env python3
"""Dump one slide's shapes, geometry and text so a patch can target them.

Usage:  python3 scripts/inspect-slide.py <deck.pptx> [slide-number ...]
        python3 scripts/inspect-slide.py <deck.pptx> --list
"""

import sys

from pptx import Presentation
from pptx.util import Emu


def inches(v):
    return None if v is None else round(Emu(v).inches, 2)


def dump(slide, number):
    print(f"===== slide {number} · {len(slide.shapes)} shapes")
    for i, sh in enumerate(slide.shapes):
        print(f"--- {i} {sh.shape_type} {sh.name!r} "
              f"L{inches(sh.left)} T{inches(sh.top)} "
              f"W{inches(sh.width)} H{inches(sh.height)}")
        if sh.shape_type == 13:
            print(f"      image {sh.image.ext} {len(sh.image.blob)} bytes "
                  f"dpi={sh.image.dpi} size={sh.image.size}")
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs)
                if text.strip():
                    sizes = [r.font.size.pt if r.font.size else None
                             for r in para.runs]
                    print(f"      {text!r} {sizes}")


def main():
    deck = Presentation(sys.argv[1])
    args = sys.argv[2:]
    if not args or args[0] == "--list":
        for i, slide in enumerate(deck.slides, 1):
            text = " ".join(sh.text_frame.text for sh in slide.shapes
                            if sh.has_text_frame).replace("\n", " ")
            print(i, text[:100])
        return
    for arg in args:
        dump(deck.slides[int(arg) - 1], int(arg))


if __name__ == "__main__":
    main()
